"""Build a deduplicated 200-person Instagram delivery payload.

FINAL rows come only from the SOP-gated database export. REVIEW rows preserve
real profile candidates found in the official-profile scrape or public search,
without representing those candidates as confirmed identities.
"""
from __future__ import annotations

import json
import re
import sqlite3
import sys
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from mp import config  # noqa: E402
from mp.pipeline import Context  # noqa: E402
from mp.stages.s8_score import build_facts  # noqa: E402
from mp.stages.s9_export import _mandatory_ok, _row_for  # noqa: E402

DISCOVERY = ROOT / "data/instagram_runs/tavily_profiles_20260907T110521Z_o0/discovery.json"
RAW = [
    ROOT / "data/instagram_runs/apify_profile_verification_20260907T125453Z/run_tcjRgvRiph6jsNP6Q.json",
    ROOT / "data/instagram_runs/apify_profile_verification_20260907T125752Z/run_9jLuUDsCjenJM6lHg.json",
]
EXISTING_DELIVERY = ROOT / "outputs/01a03f98-46d7-7503-9c57-82b11e355428/music_industry_instagram_final_20260907.xlsx"
OUTPUT = ROOT / "data/instagram_200_delivery_20260907.json"


def norm(value: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", "", (value or "").lower())


def handle_from_url(value: str | None) -> str:
    match = re.search(r"instagram\.com/([^/?#]+)/?", value or "", re.I)
    return (match.group(1).lower() if match else "").strip()


def historical_handles() -> set[str]:
    handles: set[str] = set()
    for path in ROOT.glob("outputs/**/*.xlsx"):
        if path == EXISTING_DELIVERY or path.name.startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                rows = ws.iter_rows(values_only=True)
                header = None
                for row in rows:
                    vals = [str(v).strip() if v is not None else "" for v in row]
                    if "Person Instagram URL" in vals:
                        header = vals.index("Person Instagram URL")
                        break
                if header is None:
                    continue
                for row in rows:
                    if header < len(row):
                        h = handle_from_url(str(row[header] or ""))
                        if h:
                            handles.add(h)
            wb.close()
        except Exception:
            continue
    return handles


def existing_final_rows() -> list[dict]:
    wb = openpyxl.load_workbook(EXISTING_DELIVERY, read_only=True, data_only=True)
    ws = wb["Instagram Final"]
    columns = [cell.value for cell in ws[4]]
    rows = [dict(zip(columns, row)) for row in ws.iter_rows(min_row=5, values_only=True)]
    wb.close()
    return rows


def base_rows() -> dict[int, dict]:
    ctx = Context.create(db_path="data/master.sqlite")
    sop = config.sop()
    columns = sop["sheets"]["instagram"]["columns"]
    output: dict[int, dict] = {}
    try:
        scores = {r["person_id"]: r for r in ctx.db.q(
            "SELECT * FROM scores WHERE channel='linkedin' ORDER BY total DESC"
        )}
        for person_id, score in scores.items():
            facts = build_facts(ctx, person_id)
            row = _row_for(facts, score, "linkedin", sop)
            output[person_id] = {column: row.get(column, "") for column in columns}
    finally:
        ctx.close()
    return output


def similarity(person: str, profile_name: str, handle: str) -> float:
    target = norm(person)
    name = norm(profile_name)
    h = norm(handle)
    if target and name == target:
        return 3.0
    name_ratio = SequenceMatcher(None, target, name).ratio() if name else 0
    handle_ratio = SequenceMatcher(None, target, h).ratio() if h else 0
    return max(name_ratio * 1.8, handle_ratio)


def main() -> int:
    discovery = json.loads(DISCOVERY.read_text(encoding="utf-8"))
    person_by_handle: dict[str, list[dict]] = {}
    discovery_by_person = {int(r["person_id"]): r for r in discovery}
    for row in discovery:
        for attempt in row.get("attempts", []):
            handle = (attempt.get("handle") or "").lower()
            if handle:
                person_by_handle.setdefault(handle, []).append(row)

    raw_profiles: list[dict] = []
    for path in RAW:
        raw_profiles.extend(json.loads(path.read_text(encoding="utf-8")))

    final_rows = existing_final_rows()
    final_people = {norm(r.get("Person Name")) for r in final_rows}
    final_handles = {handle_from_url(r.get("Person Instagram URL")) for r in final_rows}
    used_handles = historical_handles() | final_handles
    base = base_rows()

    best_by_person: dict[int, tuple[float, dict]] = {}
    for profile in raw_profiles:
        handle = (profile.get("username") or handle_from_url(profile.get("url"))).lower()
        if not handle or handle in used_handles:
            continue
        if not (profile.get("fullName") or profile.get("biography")):
            continue
        for person in person_by_handle.get(handle, []):
            pid = int(person["person_id"])
            score = similarity(person.get("person", ""), profile.get("fullName", ""), handle)
            current = best_by_person.get(pid)
            if current is None or score > current[0]:
                best_by_person[pid] = (score, profile)

    review_rows: list[dict] = []
    used_people = set(final_people)

    def add_review(person: dict, handle: str, evidence: str, status: str, note: str) -> bool:
        pid = int(person["person_id"])
        name_key = norm(person.get("person"))
        if not handle or handle in used_handles or name_key in used_people or pid not in base:
            return False
        row = dict(base[pid])
        row["Person Instagram URL"] = f"https://www.instagram.com/{handle}/"
        row["Instagram Profile Status"] = status
        row["Why This Person Fits"] = (
            f"{row.get('Why This Person Fits', '').rstrip()} Instagram candidate evidence: {note} "
            f"Evidence: {evidence}"
        ).strip()
        row["Verification Status"] = "RESEARCH FURTHER"
        review_rows.append(row)
        used_handles.add(handle)
        used_people.add(name_key)
        return True

    ranked = sorted(
        best_by_person.items(),
        key=lambda item: (item[1][0], int(item[1][1].get("followersCount") or 0)),
        reverse=True,
    )
    for pid, (_, profile) in ranked:
        person = discovery_by_person.get(pid)
        if not person:
            continue
        handle = (profile.get("username") or "").lower()
        displayed = profile.get("fullName") or "display name unavailable"
        bio = (profile.get("biography") or "").strip().replace("\n", " ")
        note = (
            f"Official profile scrape confirmed an available @{handle} profile displaying "
            f"'{displayed}'. Identity/current-employer relationship is not sufficiently corroborated for FINAL."
        )
        if bio:
            note += f" Public bio: {bio[:180]}"
        add_review(person, handle, f"https://www.instagram.com/{handle}/", "PROFILE EXISTS — REVIEW", note)

    # Preserve additional candidates already researched into the local evidence
    # database. VERIFIED candidates rank first; every row still remains REVIEW
    # here unless it passed the full SOP export gate and appears in FINAL.
    con = sqlite3.connect(ROOT / "data/master.sqlite")
    con.row_factory = sqlite3.Row
    db_candidates = con.execute(
        """SELECT p.person_id,p.full_name AS person,c.canonical_name AS company,
                  p.title,li.url AS linkedin_url,ig.url AS instagram_url,
                  ig.verification_status,ig.identity_score,ig.usability
             FROM people p JOIN companies c USING(company_id)
             JOIN profiles li ON li.person_id=p.person_id AND li.channel='linkedin'
             JOIN profiles ig ON ig.person_id=p.person_id AND ig.channel='instagram'
             JOIN scores sl ON sl.person_id=p.person_id AND sl.channel='linkedin'
            WHERE p.status!='REJECTED' AND p.current_verified=1
              AND c.status='QUALIFIED' AND li.verification_status='VERIFIED'
              AND sl.verdict='FINAL PASS' AND ig.url!=''
            ORDER BY CASE ig.verification_status WHEN 'VERIFIED' THEN 0 ELSE 1 END,
                     ig.identity_score DESC,p.full_name"""
    ).fetchall()
    con.close()
    for candidate in db_candidates:
        if len(review_rows) >= 196:
            break
        person = dict(candidate)
        handle = handle_from_url(candidate["instagram_url"])
        verification = candidate["verification_status"] or "UNVERIFIED"
        note = (
            f"Existing research evidence recorded @{handle} as an Instagram profile candidate "
            f"(profile verification state: {verification}; identity score: "
            f"{candidate['identity_score'] if candidate['identity_score'] is not None else 'not scored'}). "
            "It did not pass every SOP Instagram FINAL gate and requires human corroboration."
        )
        add_review(person, handle, candidate["instagram_url"], "DATABASE CANDIDATE — REVIEW", note)

    # Fill the balance with public-search-backed profile candidates. These remain
    # REVIEW because search-result existence is not identity verification.
    for person in discovery:
        if len(review_rows) >= 196:
            break
        for result in person.get("search_results", []):
            handle = handle_from_url(result.get("url"))
            if not handle:
                continue
            title = (result.get("title") or "").strip()
            snippet = (result.get("snippet") or "").strip().replace("\n", " ")
            evidence = result.get("url") or f"https://www.instagram.com/{handle}/"
            note = (
                f"Public search returned the Instagram profile @{handle} for this prospect query"
                f" ({title[:120]}). Identity/current-employer relationship is not confirmed."
            )
            if snippet:
                note += f" Search snippet: {snippet[:180]}"
            if add_review(person, handle, evidence, "PUBLIC SEARCH CANDIDATE — REVIEW", note):
                break

    if len(review_rows) < 196:
        raise RuntimeError(f"Only built {len(review_rows)} review rows; need 196")
    review_rows = review_rows[:196]
    payload = {
        "summary": {
            "total_people": len(final_rows) + len(review_rows),
            "final_people": len(final_rows),
            "review_people": len(review_rows),
            "unique_people": len({norm(r.get('Person Name')) for r in final_rows + review_rows}),
            "unique_instagram_handles": len({handle_from_url(r.get('Person Instagram URL')) for r in final_rows + review_rows}),
        },
        "columns": list(final_rows[0].keys()),
        "final_rows": final_rows,
        "review_rows": review_rows,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({**payload["summary"], "output": str(OUTPUT)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
