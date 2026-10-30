"""Bulk-verify public Instagram candidates with Apify's profile scraper.

Candidate handles come from an auditable public-web discovery file. Admission
requires an exact normalized full-name match and the SOP SS35 identity score.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sqlite3
import sys
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from mp import config  # noqa: E402
from mp.db import utcnow  # noqa: E402
from mp.normalize import instagram_handle, normalize_person_name  # noqa: E402
from mp.stages.s7_person_verify import instagram_identity_score  # noqa: E402

BASE = "https://api.apify.com/v2"
ACTOR = "apify~instagram-profile-scraper"
ROLE_HANDLE_RE = re.compile(
    r"(?:president|ceo|chief|head|director|manager|a&r|owner|founder)\s*(?:of|at|for)?\s*@([a-z0-9._]+)",
    re.I,
)


def _conflicting_current_employer(bio: str, company_ig: str) -> bool:
    """Veto profiles that advertise a senior current role at another company.

    A historical target-company mention (for example "Founder Fearless Records")
    must not outweigh a current role-bearing handle such as "President
    @thrillerrecords". Ambiguous multi-company bios stay out of FINAL.
    """
    target = instagram_handle(company_ig).lower()
    role_handles = {handle.lower() for handle in ROLE_HANDLE_RE.findall(bio or "")}
    return bool(role_handles and (not target or any(handle != target for handle in role_handles)))


def _start(token: str, handles: list[str]) -> str:
    response = requests.post(
        f"{BASE}/acts/{ACTOR}/runs",
        headers={"Authorization": f"Bearer {token}"},
        json={"usernames": handles, "includeAboutSection": False}, timeout=60,
    )
    response.raise_for_status()
    return str((response.json().get("data") or {}).get("id") or "")


def _wait(token: str, run_id: str, max_minutes: int) -> dict[str, Any]:
    started = time.monotonic()
    while time.monotonic() - started < max_minutes * 60:
        response = requests.get(
            f"{BASE}/actor-runs/{run_id}",
            headers={"Authorization": f"Bearer {token}"}, timeout=30,
        )
        response.raise_for_status()
        run = response.json().get("data") or {}
        status = str(run.get("status") or "")
        print(f"instagram_profile_run status={status}", flush=True)
        if status in {"SUCCEEDED", "FAILED", "ABORTED", "TIMED-OUT"}:
            return run
        time.sleep(20)
    raise TimeoutError("Instagram profile actor timed out")


def _items(token: str, dataset_id: str) -> list[dict[str, Any]]:
    response = requests.get(
        f"{BASE}/datasets/{dataset_id}/items",
        headers={"Authorization": f"Bearer {token}"},
        params={"format": "json", "clean": "true"}, timeout=180,
    )
    response.raise_for_status()
    body = response.json()
    return [row for row in body if isinstance(row, dict)] if isinstance(body, list) else []


def _exact_name(a: str, b: str) -> bool:
    return bool(a and b and normalize_person_name(a) == normalize_person_name(b))


def _shared_sheet_handles() -> set[str]:
    path = ROOT / "data/reference/music_industry_shared_sheet_20260907.xlsx"
    if not path.exists():
        return set()
    workbook = load_workbook(path, read_only=True, data_only=True)
    out: set[str] = set()
    for sheet in workbook.worksheets:
        values = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value or "") for value in next(values)]
        except StopIteration:
            continue
        if "Person Instagram URL" not in headers:
            continue
        index = headers.index("Person Instagram URL")
        for row in values:
            if index < len(row):
                handle = instagram_handle(str(row[index] or ""))
                if handle:
                    out.add(handle.lower())
    return out


def _previously_scraped_handles() -> set[str]:
    """Avoid spending free-tier quota on handles already fetched in prior runs."""
    out: set[str] = set()
    for path in (ROOT / "data" / "instagram_runs").glob("apify_profile_verification_*/run_*.json"):
        try:
            rows = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if not isinstance(rows, list):
            continue
        for row in rows:
            if not isinstance(row, dict):
                continue
            handle = instagram_handle(str(row.get("username") or row.get("url") or row.get("inputUrl") or ""))
            if handle:
                out.add(handle.lower())
    return out


def main(discovery_path: str, max_minutes: int, max_handles: int,
         account_index: int, handle_offset: int = 0,
         recheck_existing: bool = False) -> int:
    token = os.environ.get("APIFY_TOKEN", "").strip()
    if not token and account_index > 0:
        tokens = config.keys("apify")
        if account_index > len(tokens):
            raise SystemExit(f"Configured Apify account {account_index} does not exist")
        token = tokens[account_index - 1]
    if not token:
        raise SystemExit("APIFY_TOKEN is required")
    discovery = json.loads(Path(discovery_path).read_text(encoding="utf-8"))
    handle_rows: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    ranked_handles: list[str] = []
    # Round-robin candidate rank across people so a small budget does not spend
    # every request on several guesses for only the first few prospects.
    for attempt_index in range(12):
        for row in discovery:
            # Browser-search discovery files store exact-name candidates in
            # ``exact_results``; web-search discovery uses ``attempts``.
            attempts = (
                row.get("attempts")
                or row.get("exact_results")
                or row.get("candidates")
                or []
            )
            if attempt_index >= len(attempts):
                continue
            attempt = attempts[attempt_index]
            handle = instagram_handle(str(attempt.get("handle") or attempt.get("url") or ""))
            if handle:
                handle_rows[handle.lower()].append(row)
                ranked_handles.append(handle.lower())

    con = sqlite3.connect(ROOT / "data/master.sqlite", timeout=60)
    con.row_factory = sqlite3.Row
    existing = {str(row[0]).lower() for row in con.execute(
        "SELECT handle FROM profiles WHERE channel='instagram' AND handle IS NOT NULL"
    )} | _shared_sheet_handles() | _previously_scraped_handles()
    handles = list(dict.fromkeys(
        handle for handle in ranked_handles if recheck_existing or handle not in existing
    ))
    if handle_offset > 0:
        handles = handles[handle_offset:]
    if max_handles > 0:
        handles = handles[:max_handles]
    print(f"candidate_handles={len(handle_rows)} new_handles={len(handles)}", flush=True)
    if not handles:
        con.close()
        return 0

    run_id = _start(token, handles)
    if not run_id:
        con.close()
        raise RuntimeError("Apify did not return a run ID")
    run = _wait(token, run_id, max_minutes)
    if run.get("status") != "SUCCEEDED":
        con.close()
        raise RuntimeError(f"Instagram profile actor ended with {run.get('status')}")
    items = _items(token, str(run.get("defaultDatasetId") or ""))
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    folder = ROOT / "data/instagram_runs" / f"apify_profile_verification_{stamp}"
    folder.mkdir(parents=True, exist_ok=True)
    raw_path = folder / f"run_{run_id}.json"
    raw_path.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")

    minimum = int(config.sop()["instagram_identity_signals"]["min_score_to_verify"])
    verified: list[dict[str, Any]] = []
    for item in items:
        handle = instagram_handle(str(item.get("username") or item.get("url") or item.get("inputUrl") or ""))
        if not handle or handle.lower() not in handle_rows:
            continue
        full_name = str(item.get("fullName") or item.get("full_name") or "").strip()
        bio = str(item.get("biography") or item.get("bio") or "").strip()
        links: list[str] = []
        for value in (item.get("externalUrl"), item.get("external_url")):
            if value:
                links.append(str(value))
        for value in item.get("externalUrls") or item.get("bio_links") or []:
            if isinstance(value, dict) and (value.get("url") or value.get("lynx_url")):
                links.append(str(value.get("url") or value.get("lynx_url")))
        for candidate in handle_rows[handle.lower()]:
            person = con.execute(
                """SELECT p.person_id,p.full_name,p.current_verified,p.company_id,
                          c.canonical_name,c.instagram_url
                     FROM people p JOIN companies c USING(company_id)
                    WHERE p.person_id=? AND p.status!='REJECTED'
                      AND p.current_verified=1 AND c.status='QUALIFIED'""",
                (candidate["person_id"],),
            ).fetchone()
            if not person or not _exact_name(full_name, person["full_name"]):
                continue
            if _conflicting_current_employer(bio, person["instagram_url"] or ""):
                continue
            artists = [row[0] for row in con.execute(
                "SELECT name FROM artists WHERE company_id=? LIMIT 10", (person["company_id"],)
            )]
            score, fired = instagram_identity_score(
                bio=bio, full_name=full_name, person=person["full_name"],
                company=person["canonical_name"], external_links=links,
                company_ig=person["instagram_url"] or "", roster_artists=artists,
            )
            employer_binding = bool(
                fired.get("company_in_bio")
                or fired.get("company_tagged")
                or fired.get("external_links_confirm")
                or fired.get("roster_artists_appear")
            )
            # A same-name profile with generic words such as founder/manager is
            # insufficient. FINAL requires a signal binding the account to the
            # target music company (or its roster) in addition to identity.
            if score < minimum or not employer_binding:
                continue
            url = f"https://www.instagram.com/{handle}/"
            con.execute(
                """INSERT INTO profiles
                   (person_id,channel,url,handle,verification_status,identity_score,
                    usability,match_signals,verified_at,retrieved_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(person_id,channel) DO UPDATE SET
                     url=excluded.url,handle=excluded.handle,
                     verification_status=excluded.verification_status,
                     identity_score=excluded.identity_score,usability=excluded.usability,
                     match_signals=excluded.match_signals,verified_at=excluded.verified_at,
                     retrieved_at=excluded.retrieved_at""",
                (person["person_id"], "instagram", url, handle, "VERIFIED", score,
                 "partial" if item.get("private") or item.get("isPrivate") else "full",
                 json.dumps(fired, ensure_ascii=False, sort_keys=True), utcnow(), utcnow()),
            )
            con.execute(
                """INSERT INTO person_evidence
                   (person_id,kind,url,source,snippet,strength,retrieved_at)
                   VALUES (?,?,?,?,?,?,?)
                   ON CONFLICT(person_id,kind,url) DO UPDATE SET
                     source=excluded.source,snippet=excluded.snippet,
                     strength=excluded.strength,retrieved_at=excluded.retrieved_at""",
                (person["person_id"], "profile", url, "apify_instagram_profile",
                 f"Exact profile-name match for {person['full_name']}; SOP identity score {score}; visible bio/links: {bio[:700]}",
                 "strongest", utcnow()),
            )
            verified.append({
                "person_id": person["person_id"], "person": person["full_name"],
                "company": person["canonical_name"], "handle": handle,
                "identity_score": score, "signals": fired,
            })
            break
    con.commit()
    con.close()
    (folder / "verified.json").write_text(
        json.dumps(verified, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"dataset_items={len(items)} verified={len(verified)} raw={raw_path}")
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("discovery_path")
    parser.add_argument("--max-minutes", type=int, default=45)
    parser.add_argument("--max-handles", type=int, default=0)
    parser.add_argument("--handle-offset", type=int, default=0)
    parser.add_argument("--account-index", type=int, default=0)
    parser.add_argument("--recheck-existing", action="store_true")
    args = parser.parse_args()
    raise SystemExit(main(
        args.discovery_path, max(5, args.max_minutes), max(0, args.max_handles),
        max(0, args.account_index), max(0, args.handle_offset), args.recheck_existing,
    ))
