"""Write today's NEW people to `new data/data_<YYYY-MM-DD>.xlsx`.

"New" means the person reached a final sheet for the first time today, which is
read from `sheet_membership` in the database rather than from a run directory, so
it stays correct across restarts and re-exports. Re-running on the same day
rewrites that day's file rather than creating duplicates.
"""
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook                                    # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill         # noqa: E402
from openpyxl.utils import get_column_letter                     # noqa: E402

from mp import config                                            # noqa: E402
from mp.db import Database                                       # noqa: E402

OUT_DIR = ROOT / "new data"
NF = "Not Found"
COLUMNS = ["Channel", "Person Name", "Current Job Title", "Company Name", "Country",
           "Company Website", "Person LinkedIn URL", "Person Instagram URL",
           "Estimated Artist Roster", "Qualification Score", "Verification Status",
           "Why This Person Fits", "First Added"]


def main(day: str | None = None) -> int:
    day = day or datetime.now(timezone.utc).strftime("%Y-%m-%d")
    db = Database("data/master.sqlite")
    sop = config.sop()
    countries = config.approved_countries()
    try:
        rows = db.q("""SELECT m.channel, m.first_seen_at, p.person_id, p.full_name,
                              p.title, c.canonical_name, c.country, c.website,
                              c.roster_display, s.total, s.verdict, s.why_person
                       FROM sheet_membership m
                       JOIN people p USING(person_id)
                       JOIN companies c ON c.company_id = p.company_id
                       JOIN scores s ON s.person_id = p.person_id
                                    AND s.channel = m.channel
                       WHERE substr(m.first_seen_at, 1, 10) = ?
                         AND p.status != 'REJECTED'
                         AND c.status = 'QUALIFIED'
                         AND s.verdict = 'FINAL PASS'
                       ORDER BY s.total DESC, p.full_name""", (day,))

        out = []
        for r in rows:
            prof = {x["channel"]: x["url"] for x in db.q(
                "SELECT channel, url FROM profiles WHERE person_id=?", (r["person_id"],))}
            out.append({
                "Channel": "LinkedIn" if r["channel"] == "linkedin" else "Instagram",
                "Person Name": r["full_name"],
                "Current Job Title": r["title"] or "",
                "Company Name": r["canonical_name"],
                "Country": countries.get(r["country"], r["country"] or ""),
                "Company Website": r["website"] or NF,
                "Person LinkedIn URL": prof.get("linkedin", NF),
                "Person Instagram URL": prof.get("instagram", NF),
                "Estimated Artist Roster": r["roster_display"] or NF,
                "Qualification Score": r["total"],
                "Verification Status": r["verdict"],
                "Why This Person Fits": r["why_person"] or "",
                "First Added": r["first_seen_at"][:19].replace("T", " "),
            })

        OUT_DIR.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "data"
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in out:
            ws.append([row.get(c, "") for c in COLUMNS])
        for i, col in enumerate(COLUMNS, start=1):
            width = 52 if col.startswith("Why") else (38 if "URL" in col or "Website" in col else 22)
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        path = OUT_DIR / f"data_{day}.xlsx"
        tmp = path.with_suffix(".xlsx.tmp")
        wb.save(tmp)
        import os
        os.replace(tmp, path)          # atomic; never leaves a partial file
        print(f"written: {path}")
        print(f"  {len(out)} new people first added on {day}")
        for r in out[:10]:
            print(f"    {r['Qualification Score']}  {r['Person Name'][:22]:24s} "
                  f"{r['Current Job Title'][:28]:30s} {r['Company Name'][:24]}")
        if len(out) > 10:
            print(f"    ... and {len(out) - 10} more")
    finally:
        db.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1] if len(sys.argv) > 1 else None))
