"""Extract net-new SOP FINAL rows for the 2026-09-02 ASAP delivery."""
from __future__ import annotations

import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from mp import config  # noqa: E402
from mp.normalize import instagram_handle, linkedin_person_url  # noqa: E402
from mp.pipeline import Context  # noqa: E402
from mp.stages.s8_score import build_facts  # noqa: E402
from mp.stages.s9_export import _mandatory_ok, _row_for  # noqa: E402


OUTPUT = ROOT / "data/asap_dual_final_20260902.json"


def linkedin_key(value: object) -> str:
    return linkedin_person_url(str(value or "")).lower()


def instagram_key(value: object) -> str:
    handle = instagram_handle(str(value or ""))
    return handle.lower() if handle else ""


def workbook_history() -> tuple[set[str], set[str]]:
    linkedin: set[str] = set()
    instagram: set[str] = set()
    paths = [
        ROOT / "Music_Industry_Prospects.xlsx",
        ROOT / "data/reference/music_industry_shared_sheet_20260907.xlsx",
        *sorted((ROOT / "outputs").rglob("*.xlsx")),
    ]
    for path in paths:
        if not path.exists():
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            values = ws.iter_rows(values_only=True)
            try:
                headers = [str(x or "") for x in next(values)]
            except StopIteration:
                continue
            if "Person LinkedIn URL" in headers:
                index = headers.index("Person LinkedIn URL")
                for row in values:
                    if index < len(row) and (key := linkedin_key(row[index])):
                        linkedin.add(key)
            values = ws.iter_rows(min_row=2, values_only=True)
            if "Person Instagram URL" in headers:
                index = headers.index("Person Instagram URL")
                for row in values:
                    if index < len(row) and (key := instagram_key(row[index])):
                        instagram.add(key)
    return linkedin, instagram


def add_run_history(linkedin: set[str]) -> None:
    for path in ROOT.glob("runs/*/linkedin_final.csv"):
        with path.open(newline="", encoding="utf-8-sig") as handle:
            for row in csv.DictReader(handle):
                if key := linkedin_key(row.get("Person LinkedIn URL")):
                    linkedin.add(key)


def extract_channel(ctx: Context, channel: str, history: set[str]) -> tuple[list[str], list[dict]]:
    sop = config.sop()
    columns = sop["sheets"][channel]["columns"]
    mandatory = sop["sheets"][channel]["mandatory"]
    not_found = sop["sheets"]["not_found_literal"]
    rows_out: list[dict] = []
    seen: set[str] = set()
    per_company: dict[int, int] = defaultdict(int)
    scores = ctx.db.q(
        """SELECT s.*,p.company_id,pr.url AS profile_url
             FROM scores s JOIN people p USING(person_id)
             JOIN companies c USING(company_id)
             JOIN profiles pr ON pr.person_id=p.person_id AND pr.channel=?
            WHERE s.channel=? AND s.verdict='FINAL PASS'
              AND p.status!='REJECTED' AND p.current_verified=1
              AND c.status='QUALIFIED' AND pr.verification_status='VERIFIED'
            ORDER BY s.total DESC,c.canonical_name,p.full_name""",
        (channel, channel),
    )
    for score in scores:
        facts = build_facts(ctx, score["person_id"])
        raw_url = facts.get("linkedin_url") if channel == "linkedin" else facts.get("instagram_url")
        key = linkedin_key(raw_url) if channel == "linkedin" else instagram_key(raw_url)
        if not key or key in history or key in seen:
            continue
        company_id = int(facts["company_id"])
        if per_company[company_id] >= 10:
            continue
        row = _row_for(facts, score, channel, sop)
        ok, _ = _mandatory_ok(row, mandatory, not_found)
        if not ok:
            continue
        rows_out.append({column: row.get(column, "") for column in columns})
        seen.add(key)
        per_company[company_id] += 1
    return columns, rows_out


def main() -> int:
    old_linkedin, old_instagram = workbook_history()
    add_run_history(old_linkedin)
    ctx = Context.create(db_path="data/master.sqlite")
    try:
        li_columns, li_rows = extract_channel(ctx, "linkedin", old_linkedin)
        ig_columns, ig_rows = extract_channel(ctx, "instagram", old_instagram)
    finally:
        ctx.close()
    payload = {
        "linkedin": {"columns": li_columns, "rows": li_rows},
        "instagram": {"columns": ig_columns, "rows": ig_rows},
        "history": {"linkedin_urls": len(old_linkedin), "instagram_handles": len(old_instagram)},
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "linkedin_rows": len(li_rows), "instagram_rows": len(ig_rows),
        "linkedin_companies": len({r["Company Name"] for r in li_rows}),
        "instagram_companies": len({r["Company Name"] for r in ig_rows}),
        "output": str(OUTPUT),
    }))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
