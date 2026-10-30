"""S9 -- Final sheet generation (SS47, SS48, SS50, SS52, SS53).

Emits the two deliverables with the exact SOP column order, plus the separate
review/rejected outputs that must never be mixed into the final sheets.

Enforced here:
  SS45  only FINAL PASS (>=90 with all hard gates) enters a final sheet
  SS47/48 exact column order and per-sheet mandatory fields
  SS50  one person per row -- never a comma-joined cell
  SS52  duplicate control, strongest identifier first
  SS53  no fabricated URLs; absent optional URLs are written "Not Found"
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from .. import config
from ..normalize import person_dedupe_keys
from ..pipeline import Context

REVIEW_COLUMNS = [
    "entity_type", "entity_id", "name", "company", "status", "missing_or_unclear",
    "how_to_resolve", "evidence_urls", "suggested_verdict", "suggested_reason",
    "verdict",  # <- the human fills this in; re-ingested by `mp review-ingest`
]
REJECT_COLUMNS = ["entity_type", "name", "gate", "reason_code", "reason", "rejected_at"]


def _row_for(facts: dict, score_row, channel: str, sop: dict) -> dict:
    """Map facts onto the SS47/SS48 column names."""
    nf = sop["sheets"]["not_found_literal"]
    types = sop["company_types"]
    common = {
        "Company Name": facts["company_name"],
        "Company Website": facts["company_website"] or nf,
        "Company LinkedIn URL": facts["company_li"] or nf,
        "Company Instagram URL": facts["company_ig"] or nf,
        "Company Type": types.get(facts["company_type"], {}).get("label", ""),
        "Country": config.approved_countries().get(facts["country"], facts["country"]),
        "Estimated Artist Roster": facts["roster_display"] or nf,
        "Roster Evidence URL": facts["roster_evidence_url"] or nf,
        "Recent Activity / Intent Signal": ", ".join(
            c.replace("_", " ") for c in sorted(facts.get("signal_categories") or [])[:4]),
        "Activity Evidence URL": facts.get("activity_evidence_url") or nf,
        "Why Company Fits": score_row["why_company"] or "",
        "Person Name": facts["person_name"],
        "Current Job Title": facts["person_title"],
        "Person LinkedIn URL": facts["linkedin_url"] or nf,
        "Person Instagram URL": facts["instagram_url"] or nf,
        "Why This Person Fits": score_row["why_person"] or "",
        "Qualification Score": score_row["total"],
        "Verification Status": score_row["verdict"],
    }
    if channel == "instagram":
        common["Instagram Profile Status"] = facts.get("instagram_status", "")
    return common


def _write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _mandatory_ok(row: dict, mandatory: list[str], nf: str) -> tuple[bool, list[str]]:
    """SS47/SS48: a row missing any mandatory field cannot be written, whatever the score."""
    missing = [c for c in mandatory
               if not str(row.get(c, "")).strip() or str(row.get(c)) == nf]
    return not missing, missing


def run(ctx: Context, *, out_dir: Path | None = None) -> dict:
    from .s8_score import build_facts

    sop = config.sop()
    nf = sop["sheets"]["not_found_literal"]
    out = Path(out_dir or ctx.run_dir)
    out.mkdir(parents=True, exist_ok=True)

    sheets: dict[str, list[dict]] = {"linkedin": [], "instagram": []}
    seen: dict[str, set[tuple[str, str]]] = {"linkedin": set(), "instagram": set()}
    blocked: list[dict] = []

    rows = ctx.db.q("""SELECT s.*, p.full_name, c.canonical_name
                       FROM scores s
                       JOIN people p USING(person_id)
                       JOIN companies c ON c.company_id = p.company_id
                       -- A score row outlives a later rejection, so the person and
                       -- company status must be re-checked at export time. Without
                       -- this, someone purged as a parsing failure or an employer
                       -- mismatch still reaches the final sheet.
                       WHERE p.status != 'REJECTED'
                         AND c.status = 'QUALIFIED'
                       ORDER BY s.total DESC""")

    for score_row in rows:
        if score_row["verdict"] != "FINAL PASS":
            continue
        channel = score_row["channel"]
        facts = build_facts(ctx, score_row["person_id"])
        if not facts:
            continue

        # SS52 duplicate control -- strongest identifier first.
        keys = person_dedupe_keys(linkedin=facts["linkedin_url"],
                                  instagram=facts["instagram_url"],
                                  name=facts["person_name"],
                                  company_id=facts["company_id"],
                                  title=facts["person_title"])
        if any(k in seen[channel] for k in keys):
            ctx.bump("duplicates_suppressed")
            continue

        row = _row_for(facts, score_row, channel, sop)
        ok, missing = _mandatory_ok(row, sop["sheets"][channel]["mandatory"], nf)
        if not ok:
            blocked.append({"entity_type": "person", "entity_id": score_row["person_id"],
                            "name": facts["person_name"], "company": facts["company_name"],
                            "status": f"{channel}: FINAL PASS blocked",
                            "missing_or_unclear": ", ".join(missing),
                            "how_to_resolve": "Populate the mandatory field(s); SS53 "
                                              "forbids fabricating a value.",
                            "evidence_urls": "", "suggested_verdict": "RESEARCH_FURTHER",
                            "suggested_reason": "Mandatory sheet field missing",
                            "verdict": ""})
            continue

        seen[channel].update(keys)
        sheets[channel].append(row)

    # ---- final sheets (SS47 / SS48) ----
    for channel in ("linkedin", "instagram"):
        _write_csv(out / f"{channel}_final.csv",
                   sop["sheets"][channel]["columns"], sheets[channel])

    # ---- review queue (with optional Groq pre-adjudication) ----
    from ..sources import groq
    adjudicate = groq.available(ctx.http)
    review_rows = list(blocked)
    for row in ctx.db.q("""SELECT company_id, canonical_name, status, gate_result,
                                  country, company_type, roster_display, website
                           FROM companies WHERE status IN ('REVIEW','UNRESOLVED')"""):
        gates = json.loads(row["gate_result"] or "{}")
        unclear = [f"{k} ({v['name']})" for k, v in gates.items() if v["result"] != "PASS"]
        how = "; ".join(v["reason"] for v in gates.values()
                        if v["result"] != "PASS") or \
              "No official domain resolved. SS53: do not guess a URL."
        suggested = reason = ""
        if adjudicate:
            # Advisory only. The human `verdict` column is what the pipeline acts on;
            # Groq may judge only from facts already collected (see sources/groq.py).
            suggested, reason = groq.adjudicate_review(
                ctx.http, "company",
                {"name": row["canonical_name"], "country": row["country"],
                 "company_type": row["company_type"],
                 "roster": row["roster_display"], "website": row["website"]},
                unclear)
            ctx.bump("groq_adjudications")

        review_rows.append({
            "entity_type": "company", "entity_id": row["company_id"],
            "name": row["canonical_name"], "company": row["canonical_name"],
            "status": row["status"],
            "missing_or_unclear": ", ".join(unclear) or "domain unresolved",
            "how_to_resolve": how[:500], "evidence_urls": row["website"] or "",
            "suggested_verdict": suggested, "suggested_reason": reason,
            "verdict": ""})
    _write_csv(out / "companies_review.csv", REVIEW_COLUMNS, review_rows)

    # ---- rejections ----
    reject_rows = [{"entity_type": "company", "name": r["canonical_name"],
                    "gate": r["gate"], "reason_code": r["reason_code"],
                    "reason": r["reason_text"], "rejected_at": r["rejected_at"]}
                   for r in ctx.db.q("""SELECT c.canonical_name, r.* FROM company_rejections r
                                        JOIN companies c USING(company_id)
                                        ORDER BY r.rejected_at DESC""")]
    _write_csv(out / "rejected.csv", REJECT_COLUMNS, reject_rows)

    # ---- combined workbook ----
    _write_xlsx(out / "final_prospects.xlsx", sheets, sop)

    stats = {"linkedin_final": len(sheets["linkedin"]),
             "instagram_final": len(sheets["instagram"]),
             "review_rows": len(review_rows), "rejected_rows": len(reject_rows),
             "blocked_missing_fields": len(blocked)}
    ctx.stats.setdefault("s9", {}).update(stats)
    _write_report(ctx, out, stats)
    return stats


def _write_xlsx(path: Path, sheets: dict, sop: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True)

    for channel, title in (("linkedin", "LinkedIn Final"),
                           ("instagram", "Instagram Final")):
        ws = wb.create_sheet(title)
        columns = sop["sheets"][channel]["columns"]
        ws.append(columns)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheets[channel]:
            ws.append([row.get(c, "") for c in columns])
        for i, col in enumerate(columns, start=1):
            width = 46 if col.startswith("Why") else (34 if "URL" in col else 22)
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
        ws.freeze_panes = "A2"
    wb.save(path)


def _write_report(ctx: Context, out: Path, stats: dict) -> None:
    ledger = ctx.ledger.summary()
    status_rows = ctx.db.q("SELECT status, COUNT(*) n FROM companies GROUP BY status")
    gate_rows = ctx.db.q("""SELECT gate, COUNT(*) n FROM company_rejections
                            GROUP BY gate ORDER BY n DESC""")

    lines = [
        f"# Run report - {ctx.run_id}", "",
        f"- config hash: `{config.config_hash()}`",
        f"- SOP version: {config.sop()['sop_version']}",
        f"- countries: {', '.join(ctx.countries) or '(default)'}", "",
        "## Final output", "",
        f"- **LinkedIn Final rows:** {stats['linkedin_final']}",
        f"- **Instagram Final rows:** {stats['instagram_final']}",
        f"- review queue: {stats['review_rows']}",
        f"- rejected: {stats['rejected_rows']}",
        f"- blocked on missing mandatory fields: {stats['blocked_missing_fields']}", "",
        "## Company funnel", "",
        "| status | count |", "|---|---|",
    ]
    lines += [f"| {r['status']} | {r['n']} |" for r in status_rows]
    if gate_rows:
        lines += ["", "## Rejections by gate", "", "| gate | count |", "|---|---|"]
        lines += [f"| {r['gate']} | {r['n']} |" for r in gate_rows]

    lines += ["", "## API budget", "",
              "| provider | calls | cache hits | cache rate | cost | tripped |",
              "|---|---|---|---|---|---|"]
    for provider, s in ledger.items():
        lines.append(f"| {provider} | {s['calls']} | {s['cache_hits']} | "
                     f"{s['cache_rate']:.0%} | ${s['cost']:.4f} | "
                     f"{s['trip_reason'] or '-'} |")
    lines += ["", f"**Total run cost: ${ctx.ledger.total_cost():.4f}**", "",
              "## Stage stats", "", "```json",
              json.dumps(ctx.stats, indent=2, default=str), "```"]
    (out / "run_report.md").write_text("\n".join(lines), encoding="utf-8")
